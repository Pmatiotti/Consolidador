"""Gerador de Excel consolidado."""

import logging
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
from openpyxl.utils import get_column_letter

from models.asset import Asset

logger = logging.getLogger(__name__)

# Cores
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
ALT_FILL = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
NORMAL_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="2F5496")
SUBTITLE_FONT = Font(name="Arial", bold=True, size=11, color="2F5496")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)

COLUMNS = [
    ("Corretora", 18),
    ("Ativo", 45),
    ("Tipo Ativo", 14),
    ("Data Aplicação", 15),
    ("Indexador", 14),
    ("Taxa (%)", 12),
    ("Vencimento", 15),
    ("Liquidez", 12),
    ("Valor Aplicado (R$)", 20),
    ("Valor Bruto (R$)", 20),
    ("Valor Líquido (R$)", 20),
    ("Classe", 22),
    ("Subclasse", 22),
]


def export_excel(assets: List[Asset], output_path: str, ref_date: Optional[str] = None):
    """Exporta lista de ativos para Excel com abas Carteira e Resumo."""
    wb = Workbook()

    _create_carteira_sheet(wb, assets)
    _create_resumo_sheet(wb, assets, ref_date)

    wb.save(output_path)
    logger.info(f"Excel salvo em {output_path}")


def _create_carteira_sheet(wb: Workbook, assets: List[Asset]):
    ws = wb.active
    ws.title = "Carteira Consolidada"

    # Headers
    for col_idx, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data rows
    for row_idx, asset in enumerate(assets, 2):
        liq_display = asset.valor_liquido if asset.valor_liquido is not None else asset.valor_bruto
        values = [
            asset.corretora,
            asset.ativo,
            asset.tipo_ativo,
            asset.data_aplicacao,
            asset.indexador,
            asset.taxa,
            asset.vencimento,
            asset.liquidez,
            asset.valor_aplicado,
            asset.valor_bruto,
            liq_display,
            asset.classe,
            asset.subclasse,
        ]

        fill = ALT_FILL if row_idx % 2 == 0 else PatternFill()

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = NORMAL_FONT
            cell.border = THIN_BORDER
            cell.fill = fill

            # Formatação numérica
            if col_idx == 6 and value is not None:  # Taxa
                cell.number_format = '0.00"%"'
            elif col_idx in (9, 10, 11) and value is not None:  # Valores monetários
                cell.number_format = '#,##0.00'

    # Freeze e auto-filtro
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(assets) + 1}"


def _create_resumo_sheet(wb: Workbook, assets: List[Asset], ref_date: Optional[str] = None):
    ws = wb.create_sheet("Resumo")

    row = 1
    # Título
    title = "Resumo da Carteira Consolidada"
    if ref_date:
        title += f" — {ref_date}"
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 2

    # Tabela 1: Totais por Corretora
    row = _write_summary_table(
        ws, row, "Totais por Corretora",
        _group_totals(assets, lambda a: a.corretora), assets
    )
    row += 1

    # Tabela 2: Totais por Subclasse
    row = _write_summary_table(
        ws, row, "Totais por Subclasse",
        _group_totals(assets, lambda a: a.subclasse), assets
    )
    row += 1

    # Tabela 3: Totais por Indexador
    row = _write_summary_table(
        ws, row, "Totais por Indexador",
        _group_totals(assets, lambda a: a.indexador), assets
    )
    row += 1

    # Patrimônio total
    total_bruto = sum(a.valor_bruto for a in assets if a.valor_bruto)
    total_liquido = sum(
        (a.valor_liquido if a.valor_liquido is not None else a.valor_bruto) or 0
        for a in assets
    )

    cell = ws.cell(row=row, column=1, value="Patrimônio Total Bruto")
    cell.font = SUBTITLE_FONT
    cell = ws.cell(row=row, column=2, value=total_bruto)
    cell.number_format = '#,##0.00'
    cell.font = Font(name="Arial", bold=True, size=11)
    row += 1

    cell = ws.cell(row=row, column=1, value="Patrimônio Total Líquido")
    cell.font = SUBTITLE_FONT
    cell = ws.cell(row=row, column=2, value=total_liquido)
    cell.number_format = '#,##0.00'
    cell.font = Font(name="Arial", bold=True, size=11)

    # Ajustar larguras
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 12


def _group_totals(assets: List[Asset], key_func) -> dict:
    """Agrupa ativos por chave e soma bruto/líquido."""
    totals = {}
    for a in assets:
        key = key_func(a)
        if key not in totals:
            totals[key] = {"bruto": 0.0, "liquido": 0.0}
        if a.valor_bruto:
            totals[key]["bruto"] += a.valor_bruto
        liq = a.valor_liquido if a.valor_liquido is not None else a.valor_bruto
        if liq:
            totals[key]["liquido"] += liq
    return totals


def _write_summary_table(ws, start_row: int, title: str, totals: dict, assets: list) -> int:
    """Escreve tabela de resumo e retorna próxima linha disponível."""
    row = start_row
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SUBTITLE_FONT
    row += 1

    # Header
    headers = ["", "Bruto (R$)", "Líquido (R$)", "% Total"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
    row += 1

    total_bruto = sum(a.valor_bruto for a in assets if a.valor_bruto) or 1

    for key in sorted(totals.keys()):
        vals = totals[key]
        pct = (vals["bruto"] / total_bruto * 100) if total_bruto else 0

        fill = ALT_FILL if (row - start_row) % 2 == 0 else PatternFill()

        cell = ws.cell(row=row, column=1, value=key)
        cell.font = NORMAL_FONT
        cell.border = THIN_BORDER
        cell.fill = fill

        cell = ws.cell(row=row, column=2, value=vals["bruto"])
        cell.number_format = '#,##0.00'
        cell.font = NORMAL_FONT
        cell.border = THIN_BORDER
        cell.fill = fill

        cell = ws.cell(row=row, column=3, value=vals["liquido"])
        cell.number_format = '#,##0.00'
        cell.font = NORMAL_FONT
        cell.border = THIN_BORDER
        cell.fill = fill

        cell = ws.cell(row=row, column=4, value=round(pct, 1))
        cell.number_format = '0.0"%"'
        cell.font = NORMAL_FONT
        cell.border = THIN_BORDER
        cell.fill = fill

        row += 1

    return row
